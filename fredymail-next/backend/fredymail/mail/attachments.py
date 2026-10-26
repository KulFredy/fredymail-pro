"""Ek içerik çıkarımı (opsiyonel, varsayılan kapalı).

Ekler asla çalıştırılmaz. Yalnızca PDF/XLSX/CSV/TXT için saf-Python kütüphanelerle,
boyut ve karakter sınırı altında metin çıkarılır. Parola korumalı/bozuk dosyalar atlanır.
"""

from __future__ import annotations

import io
import logging
import warnings

from .parse import decode_bytes

log = logging.getLogger(__name__)


def extract_text(payload: bytes, filename: str, content_type: str, max_chars: int) -> tuple[str | None, str]:
    # pypdf/openpyxl bozuk dosyalarda çok sayıda zararsız uyarı basar; log gürültüsü olmasın
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        txt, status = _extract(payload, filename, content_type, max_chars)
    return (txt.replace("\x00", "") if txt else txt), status


def _extract(payload: bytes, filename: str, content_type: str, max_chars: int) -> tuple[str | None, str]:
    name = filename.lower()
    try:
        if name.endswith(".pdf") or content_type == "application/pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(payload))
            if reader.is_encrypted:
                return None, "encrypted"
            out, total = [], 0
            for page in reader.pages[:50]:
                t = page.extract_text() or ""
                out.append(t)
                total += len(t)
                if total >= max_chars:
                    break
            return "\n".join(out)[:max_chars], "ok"
        if name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            out, total = [], 0
            for ws in wb.worksheets[:10]:
                for row in ws.iter_rows(values_only=True, max_row=5000):
                    line = " | ".join(str(c) for c in row if c is not None)
                    if line:
                        out.append(line)
                        total += len(line)
                    if total >= max_chars:
                        break
            wb.close()
            return "\n".join(out)[:max_chars], "ok"
        if name.endswith((".csv", ".txt")) or content_type in ("text/csv", "text/plain"):
            return decode_bytes(payload, None)[:max_chars], "ok"
    except Exception as e:
        log.info("ek metni çıkarılamadı tür=%s hata=%s", content_type, type(e).__name__)
        return None, "error"
    return None, "unsupported"
